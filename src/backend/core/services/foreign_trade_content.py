"""Product content workspace: bounded parsing, owner isolation, no silent fact promotion."""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from uuid import uuid4
import csv
import json
import re
import subprocess
import sys
import unicodedata
import warnings
import zipfile

from sqlalchemy.orm import Session
from core.db.models import Artifact, ForeignTradeContentPack
from core.services.artifact_service import store_bytes_as_artifact
from core.storage import get_storage

MAX_UPLOAD = 8 * 1024 * 1024
MAX_TOTAL = 32 * 1024 * 1024
MAX_FILES = 10
MAX_TEXT = 60000
MAX_RAW_TEXT = 100000
MAX_PIXELS = 20_000_000


class ContentPackError(Exception):
    def __init__(self, status: int, message: str):
        self.status, self.message = status, message
        super().__init__(message)


def _now():
    return datetime.now(timezone.utc).isoformat()


def safe_filename(value: str) -> str:
    name = unicodedata.normalize("NFKC", value or "").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w. -]", "_", name).strip(" .")
    if not name or "." not in name:
        raise ContentPackError(422, "文件名缺少受支持的扩展名")
    stem, ext = name.rsplit(".", 1)
    return stem[:100] + "." + ext.lower()


def _fact_rows(rows: list[dict]) -> list[dict]:
    result = []
    for row in rows[:100]:
        label = row.get("label") or row.get("label_zh") or row.get("key")
        value = row.get("value") or row.get("value_zh")
        if label is None or value is None:
            continue
        label, value = str(label).strip(), str(value).strip()
        if not label or not value:
            continue
        source = str(row.get("source_url") or "").strip()
        if source and not re.match(r"https?://", source, re.I):
            source = ""
        result.append({"label": label[:160], "value": value[:1500], "source_url": source[:2000]})
    return result


def _csv_facts(text: str) -> list[dict]:
    # Reading plain strings never evaluates CSV formulas.
    try:
        reader = csv.DictReader(StringIO(text))
        return _fact_rows([row for index, row in enumerate(reader) if index < 100])
    except (csv.Error, TypeError):
        return []


def _pdf_text(content: bytes) -> tuple[str, list[str]]:
    if not content.startswith(b"%PDF-"):
        raise ContentPackError(422, "PDF 文件头无效")
    # Parse untrusted PDFs in a child process with independent CPU/memory limits.
    # Input bytes go through stdin, never through a shell or executable file.
    script = """
import sys, io, json, resource
resource.setrlimit(resource.RLIMIT_AS, (512*1024*1024,512*1024*1024))
resource.setrlimit(resource.RLIMIT_CPU, (8,8))
from pypdf import PdfReader
reader=PdfReader(io.BytesIO(sys.stdin.buffer.read()), strict=True)
if reader.is_encrypted: raise ValueError('encrypted')
text=[]
for i,page in enumerate(reader.pages):
    if i>=30: break
    text.append((page.extract_text() or '')[:60000])
    if sum(map(len,text))>=60000: break
print(json.dumps({'text':'\\n'.join(text)[:60000], 'limited':len(reader.pages)>30 or sum(map(len,text))>=60000}))
"""
    try:
        parsed = subprocess.run([sys.executable, "-c", script], input=content,
                                stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                timeout=12, check=True)
        data = json.loads(parsed.stdout)
    except (subprocess.SubprocessError, OSError, ValueError):
        raise ContentPackError(422, "PDF 无法安全解析、已加密或超出解析限额；请改用文本资料") from None
    notices = []
    if data["limited"]:
        notices.append("PDF 仅提取前 30 页或前 60000 字，内容已截断。")
    if not data["text"].strip():
        notices.append("该 PDF 未提取到文本，可能为扫描件；本功能未执行 OCR，请手工补充事实。")
    return data["text"], notices


def _xlsx_text(content: bytes) -> tuple[str, list[dict], list[str]]:
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            names = [info.filename for info in infos]
            if (len(infos) > 256 or len(set(names)) != len(names) or
                sum(info.file_size for info in infos) > 32 * 1024 * 1024 or
                any(info.file_size > 16 * 1024 * 1024 or info.flag_bits & 1 or
                    info.file_size > max(1, info.compress_size) * 1000 or
                    ".." in PurePosixPath(info.filename).parts or info.filename.startswith("/")
                    for info in infos) or "xl/workbook.xml" not in names or
                "[Content_Types].xml" not in names or
                any("vbaproject" in name.lower() or "externallinks" in name.lower() for name in names)):
                raise ContentPackError(422, "Excel 文件包含宏、外部链接或超出安全解压限额")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        chunks, facts, notices = [], [], []
        if len(workbook.worksheets) > 10:
            notices.append("Excel 仅提取前 10 张工作表。")
        for sheet in workbook.worksheets[:10]:
            rows = []
            for row in sheet.iter_rows(max_row=min(sheet.max_row or 1000, 1000),
                                       max_col=min(sheet.max_column or 20, 20)):
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        values.append("[公式未求值]")
                        if "Excel 公式未计算，请人工核对。" not in notices:
                            notices.append("Excel 公式未计算，请人工核对。")
                    else:
                        values.append(str(cell.value)[:2000] if cell.value is not None else "")
                rows.append(values)
            output = StringIO()
            writer = csv.writer(output)
            writer.writerows(rows)
            chunks.append("[" + sheet.title + "]\n" + output.getvalue())
            if rows:
                facts.extend(_fact_rows([dict(zip(rows[0], row)) for row in rows[1:]]))
            if (sheet.max_row or 0) > 1000 or (sheet.max_column or 0) > 20:
                notices.append("Excel 每表最多提取 1000 行、20 列，部分内容已截断。")
        workbook.close()
        return "\n".join(chunks), facts[:100], notices
    except ContentPackError:
        raise
    except Exception:
        raise ContentPackError(422, "Excel 文件损坏或格式不受支持") from None


def parse_upload(filename: str, content: bytes) -> dict:
    if not content:
        raise ContentPackError(422, "文件内容为空")
    if len(content) > MAX_UPLOAD:
        raise ContentPackError(413, "单个文件不能超过 8 MiB")
    filename = safe_filename(filename)
    extension = filename.rsplit(".", 1)[-1]
    result = {"name": filename, "content": content, "kind": "document", "text": "",
              "suggested_facts": [], "warnings": []}
    if extension in {"png", "jpg", "jpeg", "webp"}:
        from PIL import Image, ImageOps, UnidentifiedImageError
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(content)) as image:
                    if image.format not in {"PNG", "JPEG", "WEBP"} or image.width * image.height > MAX_PIXELS:
                        raise ContentPackError(422, "图片格式无效或超过 2000 万像素")
                    image.load()
                    image = ImageOps.exif_transpose(image)
                    mode = "RGBA" if "A" in image.getbands() else "RGB"
                    image = image.convert(mode)
                    if max(image.size) > 2400:
                        image.thumbnail((2400, 2400), Image.Resampling.LANCZOS)
                        result["warnings"].append("图片最长边已缩放至 2400 像素，用于安全排版。")
                    normalized = Image.new(mode, image.size)
                    normalized.paste(image)
                    target = BytesIO()
                    normalized.save(target, format="PNG")
                    payload = target.getvalue()
                    if len(payload) > MAX_TOTAL // 2:
                        raise ContentPackError(413, "规范化图片超过 16 MiB，请缩小图片")
                    result.update(content=payload, kind="image", mime_type="image/png",
                                  name=filename.rsplit(".", 1)[0] + ".png",
                                  width=normalized.width, height=normalized.height)
        except ContentPackError:
            raise
        except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            raise ContentPackError(422, "图片无法解码或超过安全像素限制") from None
        result["warnings"].append("图片仅用于素材排版，已移除元数据；未进行 OCR 或识别产品事实。")
        return result
    if extension in {"txt", "csv", "json"}:
        try:
            text = content.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("binary")
        except (UnicodeError, ValueError):
            raise ContentPackError(422, "文本资料必须为 UTF-8 编码，不能包含二进制内容") from None
        result["mime_type"] = {"txt": "text/plain", "csv": "text/csv", "json": "application/json"}[extension]
        if extension == "csv":
            result["suggested_facts"] = _csv_facts(text[:MAX_TEXT])
        elif extension == "json":
            try:
                parsed = json.loads(text)
                rows = parsed.get("facts", []) if isinstance(parsed, dict) else parsed
                if isinstance(rows, list):
                    result["suggested_facts"] = _fact_rows([row for row in rows if isinstance(row, dict)])
            except (ValueError, RecursionError):
                raise ContentPackError(422, "JSON 资料格式无效或嵌套过深") from None
        result["text"] = text
    elif extension == "xlsx":
        result["text"], result["suggested_facts"], notices = _xlsx_text(content)
        result["warnings"].extend(notices)
        result["mime_type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif extension == "pdf":
        result["text"], notices = _pdf_text(content)
        result["warnings"].extend(notices)
        result["mime_type"] = "application/pdf"
    else:
        raise ContentPackError(422, "仅支持 TXT、CSV、JSON、XLSX、PDF 和 PNG/JPEG/WebP 图片")
    if len(result["text"]) > MAX_TEXT:
        result["text"] = result["text"][:MAX_TEXT]
        result["warnings"].append("提取文本超过 60000 字，已截断；请核对原文件。")
    if result["suggested_facts"]:
        result["warnings"].append("表格事实仅作为建议，须核对并人工确认后才能生成资料包。")
    return result


class ContentPackService:
    def __init__(self, db: Session):
        self.db = db

    def _row(self, user_id: str, pack_id: str):
        row = self.db.query(ForeignTradeContentPack).filter_by(
            owner_user_id=str(user_id), pack_id=pack_id,
        ).first()
        if row is None:
            raise ContentPackError(404, "商品资料包不存在或无权访问")
        return row

    @staticmethod
    def _serialize(row):
        result = deepcopy(row.payload)
        result.update(pack_id=row.pack_id, revision=row.revision,
                      created_at=row.created_at, updated_at=row.updated_at)
        return result

    def get(self, user_id: str, pack_id: str) -> dict:
        return self._serialize(self._row(user_id, pack_id))

    def list(self, user_id: str, limit: int = 100) -> list[dict]:
        rows = self.db.query(ForeignTradeContentPack).filter_by(owner_user_id=str(user_id)).order_by(
            ForeignTradeContentPack.updated_at.desc(), ForeignTradeContentPack.pack_id,
        ).limit(limit).all()
        return [self._serialize(row) for row in rows]

    def _load(self, user_id: str, pack_id: str, revision: int) -> dict:
        record = self.get(user_id, pack_id)
        if record["revision"] != revision:
            raise ContentPackError(409, "资料包已被更新，请刷新后重试")
        return record

    def _save(self, user_id: str, pack_id: str, revision: int, record: dict) -> dict:
        payload = deepcopy(record)
        for key in ("pack_id", "revision", "created_at", "updated_at"):
            payload.pop(key, None)
        count = self.db.query(ForeignTradeContentPack).filter_by(
            owner_user_id=str(user_id), pack_id=pack_id, revision=revision,
        ).update({"payload": payload, "revision": revision + 1, "updated_at": _now()},
                 synchronize_session=False)
        if count != 1:
            self.db.rollback()
            raise ContentPackError(409, "资料包已被更新，请刷新后重试")
        self.db.commit()
        self.db.expire_all()
        return self.get(user_id, pack_id)

    def create(self, user_id: str, title: str) -> dict:
        timestamp, identifier = _now(), "pack_" + uuid4().hex
        payload = {"title": title.strip() or "新品资料包", "product_name": "", "brand": "", "model": "",
                   "source_url": "", "raw_text": "", "facts": [], "suggested_facts": [],
                   "image_credit": "", "documents": [], "outputs": None, "warnings": []}
        self.db.add(ForeignTradeContentPack(pack_id=identifier, owner_user_id=str(user_id), revision=1,
                                           payload=payload, created_at=timestamp, updated_at=timestamp))
        self.db.commit()
        return self.get(user_id, identifier)

    def update(self, user_id: str, pack_id: str, revision: int, values: dict) -> dict:
        record = self._load(user_id, pack_id, revision)
        allowed = {"title", "product_name", "brand", "model", "source_url", "raw_text", "facts", "image_credit"}
        if set(values) - allowed:
            raise ContentPackError(422, "存在不支持的资料字段")
        if any(record.get(key) != value for key, value in values.items()):
            record["outputs"] = None
        record.update(deepcopy(values))
        return self._save(user_id, pack_id, revision, record)

    def upload(self, user_id: str, pack_id: str, revision: int, filename: str, content: bytes) -> dict:
        import hashlib
        record = self._load(user_id, pack_id, revision)
        if len(record["documents"]) >= MAX_FILES:
            raise ContentPackError(413, "每个资料包最多上传 10 个文件")
        parsed = parse_upload(filename, content)
        if sum(doc["size"] for doc in record["documents"]) + len(parsed["content"]) > MAX_TOTAL:
            raise ContentPackError(413, "资料包累计文件大小不能超过 32 MiB")
        artifact = store_bytes_as_artifact(
            self.db, user_id=str(user_id), content=parsed["content"], filename=parsed["name"],
            mime_type=parsed["mime_type"], parsed_text=parsed["text"], source="foreign_trade_product_content",
            extra={"pack_id": pack_id, "kind": parsed["kind"]},
        )
        doc = {"file_id": artifact.artifact_id, "name": parsed["name"], "kind": parsed["kind"],
               "mime_type": parsed["mime_type"], "size": len(parsed["content"]),
               "sha256": hashlib.sha256(parsed["content"]).hexdigest(),
               "url": f"/api/v1/foreign-trade/content-packs/{pack_id}/files/{artifact.artifact_id}",
               "text": parsed["text"], "uploaded_at": _now()}
        for key in ("width", "height"):
            if key in parsed:
                doc[key] = parsed[key]
        record["documents"].append(doc)
        if parsed["text"]:
            combined = record["raw_text"] + "\n\n[" + parsed["name"] + "]\n" + parsed["text"]
            if len(combined) > MAX_RAW_TEXT:
                parsed["warnings"].append("资料包原文累计超过 100000 字，后续内容已截断；原文件仍保留。")
            record["raw_text"] = combined[:MAX_RAW_TEXT]
        # Import suggestions stay separate until the operator confirms via PATCH facts.
        record["suggested_facts"] = (record["suggested_facts"] + parsed["suggested_facts"])[:100]
        record["warnings"] = list(dict.fromkeys(record["warnings"] + parsed["warnings"]))
        record["outputs"] = None
        try:
            return self._save(user_id, pack_id, revision, record)
        except Exception:
            # Only this operation's newly-created artifact may be cleaned up after a CAS loss.
            self.db.rollback()
            created = self.db.query(Artifact).filter_by(artifact_id=artifact.artifact_id,
                                                        user_id=str(user_id)).first()
            if created is not None:
                try:
                    get_storage().delete(created.storage_key)
                    self.db.delete(created)
                    self.db.commit()
                except Exception:
                    self.db.rollback()
            raise

    def read_file(self, user_id: str, pack_id: str, file_id: str) -> tuple[bytes, str, str]:
        import hashlib
        record = self.get(user_id, pack_id)
        doc = next((item for item in record["documents"] if item["file_id"] == file_id), None)
        if doc is None:
            raise ContentPackError(404, "文件不属于当前资料包")
        artifact = self.db.query(Artifact).filter_by(artifact_id=file_id, user_id=str(user_id)).first()
        if artifact is None or artifact.deleted_at is not None:
            raise ContentPackError(404, "文件不存在或无权访问")
        try:
            content = get_storage().download_bytes(artifact.storage_key)
        except Exception:
            raise ContentPackError(404, "原始文件暂不可用，请重新上传") from None
        if hashlib.sha256(content).hexdigest() != doc["sha256"]:
            raise ContentPackError(409, "原文件已在其他位置改变，请重新上传并确认资料")
        return content, doc["mime_type"], doc["name"]

    def _image(self, user_id: str, record: dict) -> bytes | None:
        image = next((doc for doc in reversed(record["documents"]) if doc["kind"] == "image"), None)
        return self.read_file(user_id, record["pack_id"], image["file_id"])[0] if image else None

    async def generate(self, user_id: str, pack_id: str, revision: int, confirmed: bool) -> dict:
        from core.services.foreign_trade_content_generation import generate_pack
        record = self._load(user_id, pack_id, revision)
        if confirmed is not True:
            raise ContentPackError(422, "请先人工核对并确认产品事实")
        if not record["product_name"].strip() or not record["facts"] or len(record["facts"]) > 20:
            raise ContentPackError(422, "请填写产品名称并确认 1–20 条有依据的产品事实")
        if any(not fact["label"].strip() or not fact["value"].strip() for fact in record["facts"]):
            raise ContentPackError(422, "确认的产品事实不能留空")
        image = self._image(user_id, record)
        if len(record["image_credit"]) > 600:
            raise ContentPackError(422, "图片署名不能超过 600 字，请保留作者、来源和许可证核心信息")
        if image and not record["image_credit"].strip():
            raise ContentPackError(422, "请填写图片来源或授权说明后生成")
        result = await generate_pack(deepcopy(record), image)
        if not isinstance(result, dict) or not isinstance(result.get("listing"), dict):
            raise ContentPackError(502, "生成结果格式无效，请重试")
        record["outputs"] = deepcopy(result)
        record["outputs"].update(source_revision=revision, not_published=True)
        return self._save(user_id, pack_id, revision, record)

    def download(self, user_id: str, pack_id: str, kind: str, revision: int | None = None):
        from core.services.foreign_trade_content_generation import export_pack
        record = self.get(user_id, pack_id)
        if revision is not None and record["revision"] != revision:
            raise ContentPackError(409, "资料包已更新，请从当前版本重新下载")
        if not record.get("outputs"):
            raise ContentPackError(409, "资料已变更或尚未生成，请先确认事实并生成新版")
        if kind not in {"zip", "listing", "faq", "square", "landscape", "spec"}:
            raise ContentPackError(404, "不支持的下载类型")
        image = self._image(user_id, record)
        if kind in {"square", "landscape", "spec"} and image is None:
            raise ContentPackError(422, "未上传产品原图，无法导出图片素材")
        try:
            return export_pack(record, kind, image)
        except ContentPackError:
            raise
        except ValueError:
            raise ContentPackError(422, "当前资料无法导出所选格式") from None


def demo_file(kind: str) -> tuple[bytes, str, str]:
    from importlib.resources import files
    fixtures = {"txt": ("product_facts.txt", "text/plain"),
                "csv": ("product_facts.csv", "text/csv"),
                "photo": ("raspberry-pi-5-original.jpg", "image/jpeg")}
    if kind not in fixtures:
        raise ContentPackError(404, "示例文件不存在")
    filename, mime = fixtures[kind]
    resource = files("foreign_trade").joinpath("demo_product", filename)
    try:
        return resource.read_bytes(), mime, filename
    except (FileNotFoundError, OSError):
        raise ContentPackError(404, "示例资料未随当前安装提供") from None
